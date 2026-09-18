# SPDX-License-Identifier: MIT
"""Regression tests for the scientific core and bundled equation library."""

from __future__ import annotations

import ast
import json
import math
from pathlib import Path
import re
from tempfile import TemporaryDirectory
import unittest

import numpy as np
from openpyxl import load_workbook

from carbon_calculator.calculations import (
    RangeViolation,
    area_normalized_carbon_density,
    calculate_carbon,
    calculate_site_carbon_metrics,
    project_future_carbon,
)
from carbon_calculator import data as data_module
from carbon_calculator.data import (
    RESTORATION_ENVIRONMENTS,
    SHRUB_SPECIES,
    TREE_SPECIES,
    shrub_species_for_env,
    tree_species_for_env,
)
from carbon_calculator.data2 import (
    DOMESTIC_SPECIES,
    FOREIGN_SPECIES,
    species_map as data2_species_map,
)
from carbon_calculator import species_library
from carbon_calculator.equation_eval import EvaluationError, evaluate
from carbon_calculator.excel_export import export_all_regions_to_excel
from carbon_calculator.i18n import missing_scientific_names, tr
from carbon_calculator.input_limits import (
    SHRUB_PLANTING_AREA_M2_PER_INDIVIDUAL,
    TREE_PLANTING_AREA_M2_PER_INDIVIDUAL,
    planting_area_budget,
)
from carbon_calculator.tree_simulation.growth_models import render_states
from carbon_calculator.tree_simulation.models import VisualizationInputGroup
from carbon_calculator.tree_simulation.snapshot import build_snapshot
from carbon_calculator.translations import EN
from carbon_calculator.version import __version__


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]


class LibraryTests(unittest.TestCase):
    def test_korean_ui_literals_have_english_translations(self) -> None:
        missing: set[str] = set()
        checked = 0
        for path in (REPOSITORY_ROOT / "carbon_calculator").rglob("*.py"):
            tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
            for node in ast.walk(tree):
                if not isinstance(node, ast.Call) or not node.args:
                    continue
                if not isinstance(node.func, ast.Name) or node.func.id != "tr":
                    continue
                key = node.args[0]
                if not isinstance(key, ast.Constant) or not isinstance(key.value, str):
                    continue
                if not re.search(r"[가-힣]", key.value):
                    continue
                checked += 1
                if key.value not in EN:
                    missing.add(key.value)
        self.assertGreater(checked, 200)
        self.assertEqual(sorted(missing), [])

    def test_all_bundled_species_have_english_scientific_names(self) -> None:
        names = {
            *TREE_SPECIES,
            *SHRUB_SPECIES,
            *DOMESTIC_SPECIES,
            *FOREIGN_SPECIES,
        }
        self.assertEqual(missing_scientific_names(names), [])

    def test_plot_labels_are_translated_at_call_time(self) -> None:
        source = (REPOSITORY_ROOT / "carbon_calculator" / "plotting.py").read_text(
            encoding="utf-8"
        )
        tree = ast.parse(source)
        for node in ast.walk(tree):
            if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            defaults = [*node.args.defaults, *[d for d in node.args.kw_defaults if d]]
            for default in defaults:
                with self.subTest(function=node.name):
                    self.assertFalse(
                        isinstance(default, ast.Call)
                        and isinstance(default.func, ast.Name)
                        and default.func.id == "tr",
                        "Translated defaults must be resolved when the function is called",
                    )

    def test_release_version_is_synchronised(self) -> None:
        cff = (REPOSITORY_ROOT / "CITATION.cff").read_text(encoding="utf-8")
        installer = (REPOSITORY_ROOT / "installer.iss").read_text(encoding="utf-8")
        release_notes = (REPOSITORY_ROOT / "RELEASE_NOTES.md").read_text(encoding="utf-8")
        self.assertIn(f"version: {__version__}", cff)
        self.assertIn(f'#define MyAppVersion "{__version__}"', installer)
        self.assertIn(f"# FORECAST-SW v{__version__}", release_notes)

    def test_build_verification_records_the_current_suite_size(self) -> None:
        """BUILD_VERIFICATION.md must state the number of tests this suite actually has.

        The recorded count drifted once before (18 documented against a larger
        suite), so the document is checked against live discovery rather than a
        hard-coded number.
        """
        suite = unittest.defaultTestLoader.discover(
            start_dir=str(REPOSITORY_ROOT / "tests"),
            top_level_dir=str(REPOSITORY_ROOT),
        )
        total = suite.countTestCases()
        verification = (REPOSITORY_ROOT / "BUILD_VERIFICATION.md").read_text(encoding="utf-8")
        documented = re.search(r"(\d+) tests passed", verification)
        self.assertIsNotNone(documented, "BUILD_VERIFICATION.md must record '<n> tests passed'")
        self.assertEqual(
            int(documented.group(1)),
            total,
            "BUILD_VERIFICATION.md records "
            f"{documented.group(1)} tests but the suite now has {total}; "
            "update the document after adding or removing tests",
        )

    def test_english_readme_is_served_from_both_paths(self) -> None:
        """README.md and README.en.md must hold the same English document.

        README.md is the repository landing page and README.en.md is the path
        cited as the English manual in the article metadata (C8), so the two
        are kept byte-identical apart from line endings. README.ko.md carries
        the Korean translation and is linked from both.
        """
        landing = (REPOSITORY_ROOT / "README.md").read_text(encoding="utf-8")
        english = (REPOSITORY_ROOT / "README.en.md").read_text(encoding="utf-8")
        korean = (REPOSITORY_ROOT / "README.ko.md").read_text(encoding="utf-8")

        self.assertEqual(
            landing.replace("\r\n", "\n"),
            english.replace("\r\n", "\n"),
            "README.md and README.en.md have diverged; update both together",
        )
        for name, text in (("README.md", landing), ("README.ko.md", korean)):
            with self.subTest(document=name):
                self.assertIn("](README.ko.md)", text)
        self.assertIn("[English](README.md)", korean)

    def test_release_library_counts(self) -> None:
        self.assertEqual(len(TREE_SPECIES), 7)
        self.assertEqual(len(SHRUB_SPECIES), 15)
        self.assertEqual(len(DOMESTIC_SPECIES), 30)
        self.assertEqual(len(FOREIGN_SPECIES), 25)


class SpeciesLibraryTests(unittest.TestCase):
    """평가 화면이 쓰는 통합 라이브러리 (77종) 계층."""

    def test_every_library_record_is_selectable_in_one_input_tab(self) -> None:
        """77개 레코드 전체가 교목/관목 탭 중 정확히 한 곳에 들어간다."""
        everything = species_library.all_records()
        self.assertEqual(len(everything), 77)

        trees = species_library.records_for_kind(species_library.KIND_TREE)
        shrubs = species_library.records_for_kind(species_library.KIND_SHRUB)
        self.assertEqual(len(trees) + len(shrubs), 77)
        self.assertEqual(set(trees) & set(shrubs), set())
        self.assertEqual(set(trees) | set(shrubs), set(everything))

    def test_growth_form_follows_the_predictor_variable(self) -> None:
        """설명변수가 RCD 인 레코드는 관목, DBH 인 레코드는 교목으로 분류된다."""
        for key, record in species_library.all_records().items():
            if key in species_library.KIND_OVERRIDES:
                continue
            with self.subTest(species=key):
                expected = (species_library.KIND_SHRUB
                            if "RCD" in record.var1_label.upper()
                            else species_library.KIND_TREE)
                self.assertEqual(record.kind, expected)

    def test_year_projection_includes_explicit_assumptions(self) -> None:
        records = species_library.all_records().values()
        self.assertEqual(sum(r.supports_year_projection for r in records), 77)
        self.assertEqual(sum(r.species_data is None for r in records), 55)

    @staticmethod
    def _representative_input(record) -> float:
        """유효범위 안의 대표 입력값.

        `곰솔(지상부, 경남)` 처럼 출처가 하한을 0 으로 준 레코드가 있는데 식에
        `ln(X)` 가 들어가면 X=0 에서 정의되지 않으므로, 하한이 0 이면 구간
        중앙을 쓴다(입력 위젯도 0 을 허용하지 않는다).
        """
        if not record.has_range:
            return 10.0
        if record.range_min > 0:
            return record.range_min
        return (record.range_min + record.range_max) / 2.0

    def test_every_record_yields_a_finite_diameter_curve(self) -> None:
        """77종 전부가 직경축 곡선을 만들 수 있다 (확장 레코드 포함)."""
        for key, record in species_library.all_records().items():
            with self.subTest(species=key):
                x = self._representative_input(record)
                var2 = record.var2_default if record.is_multivar else None

                carbon = species_library.carbon_per_individual(record, x, var2)
                self.assertTrue(math.isfinite(carbon))
                self.assertGreaterEqual(carbon, 0.0)

                axis = species_library.diameter_axis(record, x)
                curve = species_library.carbon_by_diameter(record, axis, 3, var2)
                self.assertEqual(len(axis), len(curve))
                self.assertTrue(all(value > 0 for value in axis))
                finite = [value for value in curve if math.isfinite(value)]
                self.assertGreater(len(finite), len(curve) // 2)

    def test_diameter_axis_covers_the_fitted_range_when_provided(self) -> None:
        records = species_library.all_records()
        bounded = records["백합나무(전체)"]          # 6~39 cm
        axis = species_library.diameter_axis(bounded, 20.0)
        self.assertAlmostEqual(float(axis[0]), bounded.range_min, places=9)
        self.assertAlmostEqual(float(axis[-1]), bounded.range_max, places=9)
        self.assertIn(20.0, [float(v) for v in axis])

    def test_diameter_axis_brackets_the_input_when_no_range_is_published(self) -> None:
        """유효범위가 없는 레코드는 입력값 주변만 훑고 0 을 포함하지 않는다."""
        unbounded = species_library.all_records()["까치박달(전체)"]
        self.assertFalse(unbounded.has_range)
        axis = species_library.diameter_axis(unbounded, 18.0)
        self.assertGreater(float(axis[0]), 0.0)
        self.assertLess(float(axis[0]), 18.0)
        self.assertGreater(float(axis[-1]), 18.0)

    def test_range_guard_rejects_inputs_outside_the_fitted_domain(self) -> None:
        record = species_library.all_records()["백합나무(전체)"]   # 6~39 cm
        species_library.check_range(record, 6.0)
        species_library.check_range(record, 39.0)
        for value in (5.9, 39.1):
            with self.subTest(diameter=value):
                with self.assertRaises(species_library.LibraryError):
                    species_library.check_range(record, value)

    def test_extension_carbon_matches_the_equation_evaluator(self) -> None:
        """확장 레코드의 탄소량 = 식 평가값 × 탄소전환계수 × 개체수."""
        record = species_library.all_records()["백합나무(전체)"]
        biomass = evaluate(record.equation, 20.0)
        expected = biomass * species_library.CARBON_FACTOR * 4
        self.assertAlmostEqual(
            species_library.carbon_total(record, 20.0, 4), expected, places=9,
        )

    def test_core_carbon_matches_the_published_allometric_formula(self) -> None:
        """core 레코드의 탄소량은 기존 calculate_carbon 경로와 일치한다."""
        for kind in (species_library.KIND_TREE, species_library.KIND_SHRUB):
            for key, record in species_library.records_for_kind(kind).items():
                if not record.is_core:
                    continue
                with self.subTest(species=key):
                    diameter = self._representative_input(record)
                    expected = calculate_carbon(key, record.species_data, diameter, 3)
                    self.assertAlmostEqual(
                        species_library.carbon_total(record, diameter, 3),
                        expected.carbon_kg, places=9,
                    )

    def test_bundled_json_is_parseable_and_licensed(self) -> None:
        payload = json.loads((REPOSITORY_ROOT / "species_data.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["_schema"]["license"], "KOGL Type 1 (Attribution)")
        self.assertEqual(payload["_schema"]["assessment_diameter_unit"], "cm")
        self.assertEqual(payload["_schema"]["TREE_BASE_equation_diameter_unit"], "cm")
        self.assertEqual(payload["_schema"]["SHRUB_SPECIES_equation_diameter_unit"], "mm")
        self.assertEqual(payload["_schema"]["growth_diameter_unit"], "cm/year")
        self.assertEqual(len(payload["TREE_BASE"]), 7)
        self.assertEqual(len(payload["SHRUB_SPECIES"]), 15)

class SiteCategoryTests(unittest.TestCase):
    """대상지 유형이 계수·생장량에 관여하는 두 경로.

    ① 원 자료에 대상지별 식이 있는 수종은 그 식을 쓴다.
    ② 그 외 수종은 기본식을 그대로 쓰고, 연 직경 생장량에만 보정계수를 곱한다.
    """

    def setUp(self) -> None:
        self._saved_env = {
            env: dict(sections)
            for env, sections in data_module.ENVIRONMENT_GROWTH_FACTORS.items()
        }
        self._saved_species = {
            name: dict(envs)
            for name, envs in data_module.SPECIES_GROWTH_FACTORS.items()
        }

    def tearDown(self) -> None:
        data_module.ENVIRONMENT_GROWTH_FACTORS = self._saved_env
        data_module.SPECIES_GROWTH_FACTORS = self._saved_species

    def test_published_site_equations_are_applied(self) -> None:
        """소나무는 「기초 DB 자료」 순번 1·2·3 의 대상지별 식을 그대로 쓴다."""
        expected = {
            "산불피해지 자연복원": (0.0737, 2.5735, 1.0, 15.0),
            "산불피해지 인공복원": (0.0722, 2.6044, 1.0, 22.0),
            "채석장 인공복원":     (0.1323, 2.2619, 1.0, 25.0),
        }
        for category, (a, b, dmin, dmax) in expected.items():
            with self.subTest(category=category):
                pine = tree_species_for_env(category)["소나무"]
                self.assertAlmostEqual(pine.a, a, places=9)
                self.assertAlmostEqual(pine.b, b, places=9)
                self.assertAlmostEqual(pine.diameter_min, dmin, places=9)
                self.assertAlmostEqual(pine.diameter_max, dmax, places=9)

    def test_species_without_site_equations_keep_one_equation(self) -> None:
        """대상지별 식이 없는 수종은 어느 대상지에서도 식이 같다."""
        for category in RESTORATION_ENVIRONMENTS[1:]:
            baseline = tree_species_for_env(RESTORATION_ENVIRONMENTS[0])
            current = tree_species_for_env(category)
            for name in baseline:
                if name in ("소나무",):      # 대상지별 식을 보유한 수종은 제외
                    continue
                with self.subTest(species=name, category=category):
                    self.assertEqual(
                        (current[name].a, current[name].b, current[name].cf,
                         current[name].diameter_min, current[name].diameter_max),
                        (baseline[name].a, baseline[name].b, baseline[name].cf,
                         baseline[name].diameter_min, baseline[name].diameter_max),
                    )

    def test_growth_factor_scales_increments_but_not_current_stock(self) -> None:
        """보정계수는 생장량에만 곱해지고 식·0년차 저장량은 건드리지 않는다."""
        category = RESTORATION_ENVIRONMENTS[1]
        before = tree_species_for_env(category)["곰솔"]
        stock_before = calculate_carbon("곰솔", before, 12.0, 4).carbon_kg

        data_module.ENVIRONMENT_GROWTH_FACTORS = {
            env: dict(sections)
            for env, sections in data_module.ENVIRONMENT_GROWTH_FACTORS.items()
        }
        data_module.ENVIRONMENT_GROWTH_FACTORS[category]["TREE_BASE"] = 1.5
        after = tree_species_for_env(category)["곰솔"]

        self.assertEqual((after.a, after.b, after.cf), (before.a, before.b, before.cf))
        self.assertAlmostEqual(
            calculate_carbon("곰솔", after, 12.0, 4).carbon_kg, stock_before, places=12)
        for attribute in ("growth_y10", "growth_y20", "growth_y21"):
            with self.subTest(attribute=attribute):
                self.assertAlmostEqual(getattr(after, attribute),
                                       getattr(before, attribute) * 1.5, places=12)

        # 50년 시나리오는 보정계수를 반영해 달라진다
        _years, base_curve = project_future_carbon(before, 10.0, 1, years=50)
        _years, scaled_curve = project_future_carbon(after, 10.0, 1, years=50)
        self.assertAlmostEqual(float(base_curve[0]), float(scaled_curve[0]), places=12)
        self.assertGreater(float(scaled_curve[-1]), float(base_curve[-1]))

    def test_species_factor_overrides_the_category_factor(self) -> None:
        category = RESTORATION_ENVIRONMENTS[2]
        data_module.ENVIRONMENT_GROWTH_FACTORS = {
            env: dict(sections)
            for env, sections in data_module.ENVIRONMENT_GROWTH_FACTORS.items()
        }
        data_module.ENVIRONMENT_GROWTH_FACTORS[category]["TREE_BASE"] = 1.5
        data_module.SPECIES_GROWTH_FACTORS = {"곰솔": {category: 2.0}}

        base = tree_species_for_env(RESTORATION_ENVIRONMENTS[0])
        current = tree_species_for_env(category)
        self.assertAlmostEqual(current["곰솔"].growth_y10,
                               base["곰솔"].growth_y10 * 2.0, places=12)
        self.assertAlmostEqual(current["편백"].growth_y10,
                               base["편백"].growth_y10 * 1.5, places=12)

    def test_shrubs_keep_one_equation_and_only_scale_growth(self) -> None:
        category = RESTORATION_ENVIRONMENTS[1]
        data_module.ENVIRONMENT_GROWTH_FACTORS = {
            env: dict(sections)
            for env, sections in data_module.ENVIRONMENT_GROWTH_FACTORS.items()
        }
        data_module.ENVIRONMENT_GROWTH_FACTORS[category]["SHRUB_SPECIES"] = 0.8
        base = shrub_species_for_env(RESTORATION_ENVIRONMENTS[0])
        current = shrub_species_for_env(category)
        for name, spec in base.items():
            with self.subTest(species=name):
                self.assertEqual((current[name].a, current[name].b, current[name].cf),
                                 (spec.a, spec.b, spec.cf))
                self.assertAlmostEqual(current[name].growth_y10,
                                       spec.growth_y10 * 0.8, places=12)

    def test_default_factors_reproduce_the_unadjusted_library(self) -> None:
        """기본 상태(모든 계수 1.0)에서는 생장량이 기본값과 같아야 한다."""
        for category in RESTORATION_ENVIRONMENTS:
            with self.subTest(category=category):
                self.assertEqual(
                    data_module.growth_factor("곰솔", category, "TREE_BASE"), 1.0)
                self.assertEqual(
                    data_module.growth_factor("회양목", category, "SHRUB_SPECIES"), 1.0)

    def test_every_species_stores_one_record_per_site_category(self) -> None:
        """대상지 공통 '기본식' 은 없다 — 네 섹션 모두 수종마다 3개 레코드를 보유한다."""
        payload = json.loads((REPOSITORY_ROOT / "species_data.json").read_text(encoding="utf-8"))
        categories = set(payload["ENVIRONMENTS"])
        self.assertEqual(len(categories), 3)
        for section in ("TREE_BASE", "SHRUB_SPECIES",
                        "DOMESTIC_SPECIES", "FOREIGN_SPECIES"):
            for name, entry in payload[section].items():
                with self.subTest(section=section, species=name):
                    self.assertNotIn("default", entry)
                    self.assertEqual(set(entry["by_env"]), categories)

    def test_coefficient_records_carry_the_full_eight_values(self) -> None:
        """계수형 섹션의 대상지 레코드는 식 5개 + 성장량 3개를 모두 담는다."""
        payload = json.loads((REPOSITORY_ROOT / "species_data.json").read_text(encoding="utf-8"))
        for section in ("TREE_BASE", "SHRUB_SPECIES"):
            for name, entry in payload[section].items():
                for category, record in entry["by_env"].items():
                    with self.subTest(section=section, species=name, category=category):
                        self.assertEqual(len(record), 8)
                        self.assertTrue(all(isinstance(v, (int, float)) for v in record))

    def test_extension_records_resolve_per_site_category(self) -> None:
        """확장 레코드도 대상지별로 조회된다 (식이 같아도 경로가 살아 있어야 한다)."""
        for category in RESTORATION_ENVIRONMENTS:
            with self.subTest(category=category):
                domestic = data2_species_map("domestic", category)
                foreign = data2_species_map("foreign", category)
                self.assertEqual(len(domestic), 30)
                self.assertEqual(len(foreign), 25)
                self.assertTrue(domestic["백합나무(전체)"].equation)


class CalculationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.pine = TREE_SPECIES["소나무"]

    def test_allometric_formula_and_exact_count_scaling(self) -> None:
        one = calculate_carbon("소나무", self.pine, 5.0, 1)
        hundred = calculate_carbon("소나무", self.pine, 5.0, 100)
        expected = self.pine.a * 5.0 ** self.pine.b * self.pine.cf
        self.assertIsNotNone(one)
        self.assertIsNotNone(hundred)
        self.assertAlmostEqual(one.carbon_kg, expected, places=12)
        self.assertAlmostEqual(hundred.carbon_kg, one.carbon_kg * 100, places=12)

    def test_area_normalization_matches_the_controlled_site_comparison(self) -> None:
        tree = 194.146028603107
        shrub = 4.743959876150
        total = tree + shrub
        areas = (400.0, 500.0, 600.0)
        metrics = [calculate_site_carbon_metrics(tree, shrub, area) for area in areas]
        densities = [item.area_normalized_kg_m2 for item in metrics]

        self.assertEqual([round(value, 4) for value in densities], [0.4972, 0.3978, 0.3315])
        self.assertEqual([item.total_carbon_kg for item in metrics], [total] * 3)
        for value, area in zip(densities, areas):
            self.assertAlmostEqual(value * area, total, places=12)

        for invalid_area in (0.0, -1.0, math.inf, math.nan):
            with self.subTest(site_area_m2=invalid_area), self.assertRaises(ValueError):
                area_normalized_carbon_density(total, invalid_area)
        with self.assertRaises(ValueError):
            area_normalized_carbon_density(math.nan, 400.0)

    def test_combined_xlsx_exports_area_normalized_density(self) -> None:
        total = 198.889988479257
        comparison_data = [
            {
                "name": f"Profile {index}",
                "area": area,
                "env": RESTORATION_ENVIRONMENTS[0],
                "tree": 194.146028603107,
                "shrub": 4.743959876150,
                "total": total,
                "density": area_normalized_carbon_density(total, area),
            }
            for index, area in enumerate((400.0, 500.0, 600.0), start=1)
        ]

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "controlled_site_comparison.xlsx"
            export_all_regions_to_excel(str(path), [], comparison_data)
            workbook = load_workbook(path, data_only=True)
            sheet = workbook[tr("지역_비교분석")]
            self.assertIn("normalized", str(sheet.cell(2, 7).value).lower())
            self.assertEqual(
                [sheet.cell(row, 6).value for row in range(3, 6)],
                [round(total, 2)] * 3,
            )
            self.assertEqual(
                [sheet.cell(row, 7).value for row in range(3, 6)],
                [0.4972, 0.3978, 0.3315],
            )
            workbook.close()

    def test_zero_quantity_is_excluded(self) -> None:
        self.assertIsNone(calculate_carbon("소나무", self.pine, 5.0, 0))

    def test_diameter_boundaries_are_enforced(self) -> None:
        calculate_carbon("소나무", self.pine, self.pine.diameter_min, 1)
        calculate_carbon("소나무", self.pine, self.pine.diameter_max, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("소나무", self.pine, self.pine.diameter_min - 0.1, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("소나무", self.pine, self.pine.diameter_max + 0.1, 1)

    def test_projection_is_deterministic_and_matches_year_zero(self) -> None:
        current = calculate_carbon("소나무", self.pine, 5.0, 100)
        years_a, carbon_a = project_future_carbon(self.pine, 5.0, 100, years=50)
        years_b, carbon_b = project_future_carbon(self.pine, 5.0, 100, years=50)
        np.testing.assert_array_equal(years_a, years_b)
        np.testing.assert_array_equal(carbon_a, carbon_b)
        self.assertAlmostEqual(carbon_a[0], current.carbon_kg, places=12)

    def test_every_legacy_shrub_equation_accepts_rcd_in_cm_without_numerical_change(self) -> None:
        payload = json.loads((REPOSITORY_ROOT / "species_data.json").read_text(encoding="utf-8"))
        baseline = payload["ENVIRONMENTS"][0]
        for name, entry in payload["SHRUB_SPECIES"].items():
            shrub = SHRUB_SPECIES[name]
            raw = entry["by_env"][baseline]
            a, b, cf, diameter_min_mm, diameter_max_mm, *_growth = raw
            for diameter_mm in (
                diameter_min_mm,
                (diameter_min_mm + diameter_max_mm) / 2.0,
                diameter_max_mm,
            ):
                with self.subTest(species=name, diameter_mm=diameter_mm):
                    diameter_cm = diameter_mm / 10.0
                    row = calculate_carbon(name, shrub, diameter_cm, 3)
                    expected = a * diameter_mm ** b * cf * 3
                    self.assertIsNotNone(row)
                    self.assertAlmostEqual(row.carbon_kg, expected, places=12)
                    self.assertAlmostEqual(row.diameter, diameter_cm, places=12)

    def test_shrub_growth_uses_the_common_centimetre_timeline(self) -> None:
        shrub = SHRUB_SPECIES["사철나무"]
        _, carbon = project_future_carbon(shrub, 1.0, 1, years=1)
        expected_diameter_mm = (1.0 + shrub.growth_y10) * 10.0
        expected = shrub.a * expected_diameter_mm ** shrub.b * shrub.cf
        self.assertAlmostEqual(carbon[1], expected, places=12)

    def test_fractional_shrub_rcd_boundaries_are_enforced_in_cm(self) -> None:
        shrub = SHRUB_SPECIES["병꽃나무"]
        self.assertAlmostEqual(shrub.diameter_min, 0.6)
        self.assertAlmostEqual(shrub.diameter_max, 3.9)
        calculate_carbon("병꽃나무", shrub, 0.6, 1)
        calculate_carbon("병꽃나무", shrub, 3.9, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("병꽃나무", shrub, 0.59, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("병꽃나무", shrub, 3.91, 1)

    def test_shrub_visualization_uses_centimetres_for_geometry_and_carbon(self) -> None:
        shrub = SHRUB_SPECIES["병꽃나무"]
        item = VisualizationInputGroup("병꽃나무", "shrub", 1.5, 1, "cm", shrub)
        snapshot = build_snapshot(
            region_name="unit-test",
            environment="",
            area_w=5,
            area_h=5,
            inputs=(item,),
        )
        state = render_states(snapshot, 0)[0]
        current = calculate_carbon("병꽃나무", shrub, 1.5, 1)
        self.assertAlmostEqual(snapshot.groups[0].diameter_by_year[0], 1.5)
        self.assertAlmostEqual(
            snapshot.groups[0].diameter_by_year[1], 1.5 + shrub.growth_y10
        )
        self.assertAlmostEqual(snapshot.groups[0].carbon_by_year_kgc[0], current.carbon_kg)
        self.assertAlmostEqual(state.diameter_m, 0.015)

    def test_combined_planting_area_guard_is_enforced_at_the_boundary(self) -> None:
        self.assertEqual(TREE_PLANTING_AREA_M2_PER_INDIVIDUAL, 1.0)
        self.assertEqual(SHRUB_PLANTING_AREA_M2_PER_INDIVIDUAL, 0.25)

        at_limit = planting_area_budget(
            area_w=10, area_h=10, tree_quantity=50, shrub_quantity=200
        )
        self.assertIsNotNone(at_limit)
        self.assertAlmostEqual(at_limit.required_area_m2, 100.0)
        self.assertFalse(at_limit.is_exceeded)
        self.assertAlmostEqual(at_limit.excess_area_m2, 0.0)

        over_limit = planting_area_budget(
            area_w=10, area_h=10, tree_quantity=80, shrub_quantity=100
        )
        self.assertIsNotNone(over_limit)
        self.assertAlmostEqual(over_limit.tree_area_m2, 80.0)
        self.assertAlmostEqual(over_limit.shrub_area_m2, 25.0)
        self.assertAlmostEqual(over_limit.required_area_m2, 105.0)
        self.assertTrue(over_limit.is_exceeded)
        self.assertAlmostEqual(over_limit.excess_area_m2, 5.0)

        self.assertIsNone(
            planting_area_budget(
                area_w=0, area_h=10, tree_quantity=1, shrub_quantity=1
            )
        )


class EquationEvaluatorTests(unittest.TestCase):
    def test_supported_equation_forms(self) -> None:
        self.assertAlmostEqual(evaluate("Y=0.063*X^2.578", 10), 0.063 * 10 ** 2.578)
        self.assertAlmostEqual(
            evaluate("ln(Y)=2.43*ln(X)-2.28", 10),
            math.exp(2.43 * math.log(10) - 2.28),
        )
        self.assertAlmostEqual(
            evaluate("Y=exp(-4.7483+1.7395*ln(X*H))", 10, 12),
            math.exp(-4.7483 + 1.7395 * math.log(120)),
        )

    def test_entire_compatibility_library_executes(self) -> None:
        for name, record in {**DOMESTIC_SPECIES, **FOREIGN_SPECIES}.items():
            if record.has_range:
                x = (record.diameter_min + record.diameter_max) / 2
            else:
                x = 10.0
            h = record.var2_default if record.is_multivar else None
            with self.subTest(name=name):
                self.assertTrue(math.isfinite(evaluate(record.equation, x, h)))

    def test_unsafe_or_unknown_syntax_is_rejected(self) -> None:
        unsafe = (
            "Y=__import__('os').system('echo unsafe')",
            "Y=(1).__class__",
            "Y=[X][0]",
            "Y=UNKNOWN(X)",
        )
        for equation in unsafe:
            with self.subTest(equation=equation), self.assertRaises(EvaluationError):
                evaluate(equation, 10)


if __name__ == "__main__":
    unittest.main()
